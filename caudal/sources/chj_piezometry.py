"""Groundwater level trend per aquifer from the Júcar's piezometric network (CHJ).

The CHJ publishes every validated measurement of its piezometers since the 1970s as one
open download, with the groundwater body (PHJ22) each piezometer belongs to. For each
piezometer we take the median static level of each year and fit a Theil–Sen slope
(the median of the slopes between every pair of years), which one odd year or a
pumping test cannot drag. An aquifer's trend is the median slope of its piezometers.
"""
from __future__ import annotations

import io
import statistics
import zipfile
from dataclasses import dataclass
from datetime import date

from . import session

DOWNLOAD = "https://aps.chj.es/down/CSV/F71_Puntos_de_control_piezometrico.zip"
PAGE = "https://www.chj.es/es-es/medioambiente/redescontrol/Paginas/Piezometria.aspx"
POINTS_SHEET = "F71_M10_Todas.xlsx"
MEASURES_PREFIX = "F71_D7_Medida en campo"
STATIC_LEVEL = "Niv. Valid. Estática (m.s.n.m.)"

FIRST_YEAR_BACK = 25      # trend over the last 25 years…
MIN_YEARS = 10            # …from piezometers with at least 10 of them measured…
RECENT_YEARS = 3          # …one of them in the last 3, so dead piezometers do not count


@dataclass(frozen=True)
class AquiferTrend:
    code: str                 # CHJ groundwater body code, e.g. 080-133
    name: str
    slope_m_per_year: float   # median of its piezometers' slopes; negative is falling
    piezometers: int
    falling: int              # piezometers falling more than 10 cm a year
    first_year: int
    last_year: int


def theil_sen(points: list[tuple[int, float]]) -> float:
    slopes = [(y2 - y1) / (x2 - x1) for i, (x1, y1) in enumerate(points) for x2, y2 in points[i + 1:] if x2 != x1]
    return statistics.median(slopes)


def _header_index(header: tuple) -> dict[str, int]:
    # The sheets come with Latin-1 headers read as UTF-8; compare on the ASCII part only.
    ascii_ = lambda s: "".join(c for c in str(s) if c.isascii())      # noqa: E731
    return {ascii_(h): i for i, h in enumerate(header)}


def _rows(book_bytes: bytes):
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(book_bytes), read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    idx = _header_index(next(rows))
    for r in rows:
        yield idx, r


def trends(archive: bytes, today: date | None = None) -> dict[str, AquiferTrend]:
    """Trend of every aquifer with enough long-running piezometers, keyed by CHJ code."""
    today = today or date.today()
    first, recent = today.year - FIRST_YEAR_BACK, today.year - RECENT_YEARS
    zf = zipfile.ZipFile(io.BytesIO(archive))
    names = {n.rsplit("/", 1)[-1]: n for n in zf.namelist()}
    aquifer_of, aquifer_name = {}, {}
    for idx, r in _rows(zf.read(names[POINTS_SHEET])):
        code = r[idx["Cd. Masa PHJ22 Hidrog."]]
        if code:
            aquifer_of[r[idx["Cd. Punto"]]] = code
            aquifer_name[code] = r[idx["Masa Subterrnea PHJ22 Hidrog."]]
    by_year: dict[str, dict[int, list[float]]] = {}
    level_key = "".join(c for c in STATIC_LEVEL if c.isascii())
    for name in sorted(n for n in names if n.startswith(MEASURES_PREFIX)):
        for idx, r in _rows(zf.read(names[name])):
            point, when, level = r[idx["Cd. Punto"]], r[idx["Fecha"]], r[idx[level_key]]
            if point in aquifer_of and when and level is not None and when.year >= first:
                by_year.setdefault(point, {}).setdefault(when.year, []).append(float(level))
    slopes: dict[str, list[tuple[float, int, int]]] = {}
    for point, years in by_year.items():
        if len(years) < MIN_YEARS or max(years) < recent:
            continue
        series = sorted((y, statistics.median(v)) for y, v in years.items())
        slopes.setdefault(aquifer_of[point], []).append((theil_sen(series), series[0][0], series[-1][0]))
    return {code: AquiferTrend(code, aquifer_name[code], statistics.median(s for s, _, _ in v), len(v),
                               sum(s < -0.1 for s, _, _ in v), min(a for _, a, _ in v), max(b for _, _, b in v))
            for code, v in slopes.items()}


def fetch(http=None) -> bytes:
    http = http or session()
    resp = http.get(DOWNLOAD, timeout=300)
    resp.raise_for_status()
    return resp.content
